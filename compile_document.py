# -*- coding: utf-8 -*-
"""
Master compiler script to generate:
PISA_2025_澳門數據與校本課程改革專題分析報告.docx
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from report_sections_part1 import (
    build_cover_and_header,
    build_executive_summary,
    build_section_1,
    build_section_2
)
from report_sections_part2 import (
    build_section_3,
    build_section_4
)
from report_sections_part3 import (
    build_section_5,
    build_section_6,
    build_section_7,
    build_section_8
)

def create_full_report():
    print("Initializing document...")
    doc = Document()
    
    # Configure Section Margins
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Configure Header
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_head = p_head.add_run("校本課程發展專題分析報告 | PISA 2025 澳門數據與校本課程改革啟示")
    r_head.font.name = "微軟正黑體"
    r_head.font.size = Pt(8.5)
    r_head.font.color.rgb = RGBColor(160, 174, 192)
    
    # Configure Footer
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_f1 = p_foot.add_run("機密文件 · 僅供校長室、教務委員會與各學科組內部研討使用")
    r_f1.font.name = "微軟正黑體"
    r_f1.font.size = Pt(8)
    r_f1.font.color.rgb = RGBColor(160, 174, 192)
    
    # Configure Default Style Font
    normal_style = doc.styles['Normal']
    normal_style.font.name = '微軟正黑體'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(45, 55, 72)
    
    # Set EastAsian font attribute on Normal style
    rPr = normal_style.element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="微軟正黑體" w:cs="Calibri"/>')
    rPr.append(rFonts)

    print("Building Cover & Metadata...")
    build_cover_and_header(doc)
    
    print("Building Executive Summary...")
    build_executive_summary(doc)
    
    print("Building Section 1: Data Sources & Methodology...")
    build_section_1(doc)
    
    print("Building Section 2: 8 Dimensions Framework...")
    build_section_2(doc)
    
    print("Building Section 3: 5 Core Contradictions...")
    build_section_3(doc)
    
    print("Building Section 4: 3-Level Implementation & Diagnostics...")
    build_section_4(doc)
    
    print("Building Section 5: P1-S6 Vertical Progression & Learning Chasms...")
    build_section_5(doc)
    
    print("Building Section 6: Action Plan & 10-Dimension Matrix...")
    build_section_6(doc)
    
    print("Building Section 7: Reform Roadmap (2026-2028)...")
    build_section_7(doc)
    
    print("Building Section 8: References & Appendices...")
    build_section_8(doc)

    out_path = r"h:\我的雲端硬碟\11_2025PISA分析\PISA_2025_澳門數據與校本課程改革專題分析報告.docx"
    print(f"Saving document to: {out_path} ...")
    doc.save(out_path)
    print("Document successfully compiled and saved!")

if __name__ == "__main__":
    create_full_report()
