# -*- coding: utf-8 -*-
"""
Full report generator for PISA 2025 School-based Curriculum Reform Analysis.
Author: Curriculum Development Assistant Director
Target: School Leadership & Academic Panels
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# --- Color Palette Constants ---
HEX_NAVY = "1A365D"       # Primary Accent: Deep Navy
HEX_SLATE = "2B6CB0"      # Secondary Accent: Steel Blue
HEX_DARK = "2D3748"       # Charcoal Body Text
HEX_MUTED = "718096"      # Muted Gray
HEX_LIGHT_BG = "F7FAFC"   # Shading Background
HEX_ROW_ALT = "EDF2F7"    # Alternating row background
HEX_ACCENT_RED = "C53030" # Alert / Warning
HEX_ACCENT_GOLD = "D69E2E"# Gold Accent
HEX_BORDER = "CBD5E0"     # Table borders
HEX_WHITE = "FFFFFF"

RGB_NAVY = RGBColor(26, 54, 93)
RGB_SLATE = RGBColor(43, 108, 176)
RGB_DARK = RGBColor(45, 55, 72)
RGB_MUTED = RGBColor(113, 128, 150)
RGB_RED = RGBColor(197, 48, 48)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_callout_border(cell, border_color=HEX_NAVY, border_size=28):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_table_borders(table, border_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{HEX_NAVY}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout(doc, text_paragraphs, title="核心洞察與管理層啟示", alert=False):
    border_color = HEX_ACCENT_RED if alert else HEX_NAVY
    bg_color = "FFF5F5" if alert else HEX_LIGHT_BG
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_callout_border(cell, border_color, border_size=28)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run_title = p.add_run(f"【{title}】\n")
    run_title.bold = True
    run_title.font.name = "微軟正黑體"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGB_RED if alert else RGB_NAVY
    
    if isinstance(text_paragraphs, str):
        text_paragraphs = [text_paragraphs]
        
    for i, para in enumerate(text_paragraphs):
        if i == 0:
            p_text = p
        else:
            p_text = cell.add_paragraph()
            p_text.paragraph_format.space_before = Pt(2)
            p_text.paragraph_format.space_after = Pt(3)
            p_text.paragraph_format.line_spacing = 1.25
        run = p_text.add_run(para)
        run.font.name = "微軟正黑體"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGB_DARK
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "微軟正黑體"
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(14)
        run.font.color.rgb = RGB_NAVY
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        run.font.size = Pt(12)
        run.font.color.rgb = RGB_SLATE
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGB_DARK
    return p

def add_para(doc, text, bold_prefix="", space_after=5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "微軟正黑體"
        r_pre.font.size = Pt(10)
        r_pre.bold = True
        r_pre.font.color.rgb = RGB_DARK
        
    run = p.add_run(text)
    run.font.name = "微軟正黑體"
    run.font.size = Pt(10)
    run.font.color.rgb = RGB_DARK
    run.italic = italic
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.2
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "微軟正黑體"
        r_pre.font.size = Pt(9.5)
        r_pre.bold = True
        r_pre.font.color.rgb = RGB_DARK
        
    run = p.add_run(text)
    run.font.name = "微軟正黑體"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGB_DARK
    return p

def add_styled_table(doc, headers, data, col_widths=None, alignment=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)
    
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = "微軟正黑體"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg = HEX_ROW_ALT if (r_idx % 2 == 1) else HEX_WHITE
        
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            if alignment and c_idx < len(alignment):
                p.alignment = alignment[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                
            run = p.add_run(str(val))
            run.font.name = "微軟正黑體"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGB_DARK
            
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return tbl

print("Ready to construct document generator.")
